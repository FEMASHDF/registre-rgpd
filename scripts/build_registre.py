"""Compile les issues « registre » en registre RGPD + plan d'actions.

Sorties (dossier registre/) :
    REGISTRE.md        registre lisible (synthèse + fiches + plan d'actions)
    registre.csv       une ligne par activité de traitement
    plan_actions.csv   écarts détectés vs référentiel CNIL
    registre.xlsx      export Excel (Synthèse, Registre, Plan d'actions)

Statut d'une fiche selon ses labels : "valide" -> Validée, sinon "À valider".
Les issues portant le label "archive" sont exclues.

Usage : python scripts/build_registre.py            (via gh CLI)
        python scripts/build_registre.py issues.json (export local)
"""
from __future__ import annotations

import csv
import datetime as dt
import json
import subprocess
import sys

from common import ROOT, load_schema, parse_issue_body, questions

OUT_DIR = ROOT / "registre"
TODO_WORDS = ("à régulariser", "à corriger", "à formaliser", "à vérifier", "non réalisée",
              "non désigné", "à faire", "à compléter", "à mettre en place", "non vérifié",
              "– en cours")
PRIORITY_HIGH = {"hebergeur_hds", "sec_authentifier", "sec_habilitations", "aipd", "dpo_designe",
                 "contrat_art28", "transfert_hors_ue", "sec_sauvegarde"}


def load_issues() -> list[dict]:
    if len(sys.argv) > 1:
        out = open(sys.argv[1], encoding="utf-8").read()
    else:
        out = subprocess.run(["gh", "issue", "list", "--state", "all",
                              "--limit", "2000", "--json",
                              "number,title,body,state,labels,url,updatedAt"],
                             check=True, capture_output=True, text=True).stdout
    # Fiche = label « registre » OU titre « [Registre] » (si les labels n'existaient pas encore)
    return [i for i in json.loads(out)
            if any(lab["name"] == "registre" for lab in i.get("labels", []))
            or (i.get("title") or "").startswith("[Registre]")]


def write_xlsx(path, schema: dict, fiches: list[dict], today: str) -> None:
    """Export Excel mis en forme : Synthèse, Registre (1 ligne / traitement), Plan d'actions."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    qs = questions(schema)
    head_fill = PatternFill("solid", fgColor="1F4E79")
    head_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")
    high_fill = PatternFill("solid", fgColor="FCE4D6")

    def table(ws, start_row, headers, rows, widths):
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row=start_row, column=c, value=h)
            cell.fill, cell.font, cell.border = head_fill, head_font, border
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for r, row in enumerate(rows, start_row + 1):
            for c, v in enumerate(row, 1):
                cell = ws.cell(row=r, column=c, value=v)
                cell.alignment, cell.border = wrap, border
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = w
        ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
        if rows:
            ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{start_row + len(rows)}"

    def title(ws, text):
        ws["A1"] = text
        ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
        ws["A2"] = f"Registre des activités de traitement (art. 30 RGPD) – généré le {today}"
        ws["A2"].font = Font(italic=True, color="595959")

    def val(v):
        return "\n".join(f"• {x}" for x in v) if isinstance(v, list) else (v or "")

    wb = Workbook()
    ws = wb.active
    ws.title = "Synthèse"
    title(ws, "FEMAS HDF – Registre RGPD")
    table(ws, 4, ["N°", "Structure", "Traitement", "Responsable", "Statut", "Nb écarts", "Dont priorité haute", "Mise à jour", "Lien"],
          [[x["issue"]["number"], x["ans"].get("structure", ""), x["ans"].get("nom_traitement", ""),
            x["ans"].get("responsable", ""), x["statut"], len(x["gaps"]),
            sum(1 for g in x["gaps"] if g[0] == "Haute"), x["ans"].get("date_maj", ""), x["issue"]["url"]]
           for x in fiches], [6, 28, 34, 30, 12, 10, 12, 14, 40])

    ws = wb.create_sheet("Registre")
    title(ws, "Fiches de traitement")
    table(ws, 4, ["N°", "Statut"] + [f["label"] for f in qs],
          [[x["issue"]["number"], x["statut"]] + [val(x["ans"].get(f["id"])) for f in qs] for x in fiches],
          [6, 12] + [45 if f["type"] in ("checkboxes", "textarea") else 28 for f in qs])

    ws = wb.create_sheet("Plan d'actions")
    title(ws, "Plan d'actions – écarts au référentiel CNIL")
    rows = [[x["issue"]["number"], x["ans"].get("structure", ""), x["ans"].get("nom_traitement", ""),
             p, t, a, "", "", ""]
            for x in fiches for p, t, a in sorted(x["gaps"], key=lambda g: g[0] != "Haute")]
    table(ws, 4, ["N°", "Structure", "Traitement", "Priorité", "Thème", "Action", "Responsable", "Échéance", "Fait le"],
          rows, [6, 24, 28, 10, 32, 60, 18, 12, 12])
    for r in range(5, 5 + len(rows)):
        if ws.cell(row=r, column=4).value == "Haute":
            for c in range(1, 10):
                ws.cell(row=r, column=c).fill = high_fill
    wb.save(path)


def detect_gaps(schema: dict, ans: dict) -> list[tuple[str, str, str]]:
    """-> [(priorité, thème, action)]"""
    gaps = []
    for f in questions(schema):
        prio = "Haute" if f["id"] in PRIORITY_HIGH else "Normale"
        val = ans.get(f["id"])
        if f["type"] == "checkboxes" and f.get("gap"):
            for opt in f["options"]:
                if opt not in (val or []):
                    gaps.append((prio, f["label"], f"Mettre en place : {opt}"))
        elif f["type"] == "dropdown" and isinstance(val, str):
            low = val.lower()
            if any(w in low for w in TODO_WORDS) or low.startswith("non –") or low == "non":
                if f["id"] in ("transfert_hors_ue", "si_partage"):
                    continue  # « Non » est ici la bonne réponse
                gaps.append((prio, f["label"], f"Réponse « {val} » : action requise"))
        elif f["type"] == "checkboxes" and f["id"] == "information":
            if any("Aucune" in v for v in (val or [])):
                gaps.append(("Haute", f["label"], "Mettre en place l'information des patients"))

    # Cohérence AIPD / DPO (référentiel §11 : groupe + SI partagé + ≥ 10 000 patients/an)
    if ans.get("si_partage") == "Oui" and ans.get("volume_patients") == "10 000 ou plus":
        if ans.get("aipd", "").startswith("Non requise"):
            gaps.append(("Haute", "AIPD", "AIPD nécessaire (SI partagé, ≥ 10 000 patients/an)"))
        if ans.get("dpo_designe", "").startswith("Non requis"):
            gaps.append(("Haute", "DPO", "DPO nécessaire (SI partagé, ≥ 10 000 patients/an)"))
    if ans.get("volume_patients") == "Ne sais pas":
        gaps.append(("Normale", "AIPD / DPO", "Évaluer la file active annuelle (seuil 10 000 patients)"))
    return gaps


def fmt(v) -> str:
    if isinstance(v, list):
        return "<br>".join(f"• {x}" for x in v) if v else "—"
    return (v or "—").replace("\n", "<br>")


def main() -> None:
    schema = load_schema()
    qs = questions(schema)
    fiches = []
    for issue in load_issues():
        labels = {lab["name"] for lab in issue.get("labels", [])}
        if "archive" in labels:
            continue
        ans = parse_issue_body(schema, issue["body"])
        fiches.append({"issue": issue, "ans": ans,
                       "statut": "Validée" if "valide" in labels else "À valider",
                       "gaps": detect_gaps(schema, ans)})
    fiches.sort(key=lambda x: (x["ans"].get("structure", ""), x["ans"].get("nom_traitement", "")))

    OUT_DIR.mkdir(exist_ok=True)
    today = dt.date.today().strftime("%d/%m/%Y")

    # ---- CSV registre
    with open(OUT_DIR / "registre.csv", "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(["N° fiche", "Statut", "Lien"] + [f["label"] for f in qs] + ["Nb écarts"])
        for x in fiches:
            row = [x["issue"]["number"], x["statut"], x["issue"]["url"]]
            for f in qs:
                v = x["ans"].get(f["id"])
                row.append(" | ".join(v) if isinstance(v, list) else (v or ""))
            w.writerow(row + [len(x["gaps"])])

    # ---- CSV plan d'actions
    with open(OUT_DIR / "plan_actions.csv", "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(["N° fiche", "Structure", "Traitement", "Priorité", "Thème", "Action",
                    "Responsable", "Échéance", "Fait le"])
        for x in fiches:
            for prio, theme, action in x["gaps"]:
                w.writerow([x["issue"]["number"], x["ans"].get("structure", ""),
                            x["ans"].get("nom_traitement", ""), prio, theme, action, "", "", ""])

    # ---- Markdown
    md = ["# FEMAS HDF – Registre des activités de traitement",
          "",
          f"_Registre des activités de traitement (art. 30 RGPD) – généré automatiquement le {today}._",
          "_Ce document doit pouvoir être présenté à la CNIL en cas de contrôle._",
          "",
          "## Synthèse", "",
          "| N° | Structure | Traitement | Responsable | Statut | Écarts | MAJ |",
          "|---|---|---|---|---|---|---|"]
    for x in fiches:
        a, i = x["ans"], x["issue"]
        md.append(f"| [#{i['number']}]({i['url']}) | {a.get('structure', '')} | {a.get('nom_traitement', '')} "
                  f"| {a.get('responsable', '')} | {x['statut']} | {len(x['gaps'])} | {a.get('date_maj', '')} |")
    if not fiches:
        md.append("| – | – | Aucune fiche | – | – | – | – |")

    md += ["", "## Fiches de traitement", ""]
    current = None
    for x in fiches:
        a, i = x["ans"], x["issue"]
        md += [f"### #{i['number']} – {a.get('nom_traitement', i['title'])}", "",
               f"Statut : **{x['statut']}** · [voir l'issue]({i['url']})", ""]
        for f in schema["fields"]:
            if f["type"] == "section":
                current = f["label"]
                md += [f"**{current}**", "", "| Rubrique | Réponse |", "|---|---|"]
                continue
            md.append(f"| {f['label']} | {fmt(a.get(f['id']))} |")
            nxt = schema["fields"].index(f) + 1
            if nxt >= len(schema["fields"]) or schema["fields"][nxt]["type"] == "section":
                md.append("")

    md += ["## Plan d'actions (écarts au référentiel CNIL)", ""]
    any_gap = False
    for x in fiches:
        if not x["gaps"]:
            continue
        any_gap = True
        md += [f"### #{x['issue']['number']} – {x['ans'].get('structure', '')} / "
               f"{x['ans'].get('nom_traitement', '')}", "",
               "| Priorité | Thème | Action |", "|---|---|---|"]
        for prio, theme, action in sorted(x["gaps"], key=lambda g: g[0] != "Haute"):
            md.append(f"| {prio} | {theme} | {action} |")
        md.append("")
    if not any_gap:
        md.append("Aucun écart détecté.")

    (OUT_DIR / "REGISTRE.md").write_text("\n".join(md) + "\n", encoding="utf-8")
    write_xlsx(OUT_DIR / "registre.xlsx", schema, fiches, today)
    print(f"{len(fiches)} fiche(s) compilée(s), "
          f"{sum(len(x['gaps']) for x in fiches)} action(s) au plan.")


if __name__ == "__main__":
    main()
